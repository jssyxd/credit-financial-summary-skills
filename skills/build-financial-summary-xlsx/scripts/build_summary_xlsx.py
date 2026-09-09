#!/usr/bin/env python3
"""按多年 audit JSON 生成「近三年及近期财务数据简表.xlsx」。
两种模式:
  A) 克隆(推荐, 样式 1:1):  --template <范例.xlsx>   载入范例后仅覆写数据格(含格式保留)
  B) 重建: 无模板时按 row_defs 样式快照从零建表(仿宋/行高/列宽/居中/换行)
约定(与范例一致): 万元=round(元/1e4,2); 0或缺失→'-'; 金额|x|>=1000→'#,##0.00'否则General;
变化率=(后-前)/|前| round4 格式0.00%; 前0后正→1, 前0后负→-1, 皆0/缺→'-'(用显示两位万计算)。
用法:
python build_summary_xlsx.py \
  --source 2023=audit2023.json --source 2024=audit2024.json \
  --source 2025=audit2025.json --source 2606=audit2606.json \
  --out 本部财务数据简表.xlsx [--template 范例.xlsx] [--period-label 2026年6月]
"""
import argparse, json, os, re, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from openpyxl.styles import Font, Alignment
from row_defs import ROWS, LONGINV, ROWHEIGHTS, COLWIDTHS, FONT_NAME, norm_key, val, collect

# 模板行号 → (B列显示文本, 序号A值或None); 51 行固定(范例 4..54)
ROWSHEET = []
_seq = {4: 1, 5: 1.1, 14: 1.2, 15: 1.3, 16: 1.4, 17: 1.5, 18: 1.6, 19: 1.7, 20: 1.8, 21: 2, 22: 2.1,
        30: 2.2, 34: 3, 38: 4, 39: 5, 40: 6, 41: 7, 42: 8, 43: 9, 44: 10, 45: 11, 46: 12, 47: 13,
        48: 14, 49: 15, 50: 16, 51: 17, 52: 18, 53: 19, 54: 20}
# 由 ROWS 顺序铺到 4..54 (ROWS 长度应=51)
for i, (label, typ, cands) in enumerate(ROWS):
    r = 4 + i
    ROWSHEET.append((r, _seq.get(r), label, typ, cands))


def stmt(doc, t):
    return {'bs': doc['bs']['current'], 'is': doc['is']['current'], 'cf': doc['cf']['current']}[t]


def amount(doc, t, cs):
    d = stmt(doc, t)
    if cs is None:
        return collect(d, *LONGINV)
    return val(d, *cs)


def wan(x):
    return None if x is None else round(x / 1e4, 2)


def ratio(nw, od):
    # 范例约定: (后-前)/|前|; 前0/缺→按方向±1; 用两位万显示值计算(避免OCR分币噪声爆炸)
    if od in (None, 0) and nw in (None, 0):
        return None
    if od in (None, 0):
        return 1.0 if (nw or 0) > 0 else (-1.0 if (nw or 0) < 0 else 0.0)
    n = nw or 0.0
    return round((n - od) / abs(od), 4)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--source', action='append', required=True,
                    help='形如 2023=audit2023.json (可多次); 顺序即 C/D/F/H 列顺序')
    ap.add_argument('--out', required=True)
    ap.add_argument('--template', default=None, help='范例xlsx; 提供则克隆样式')
    ap.add_argument('--period-label', default='2026年6月', help='近一期列表头文本')
    args = ap.parse_args()

    cols = ['C', 'D', 'F', 'H']          # 简表数据列(范例: C2023 D2024 F2025 H近一期)
    growth = [('E', 0, 1), ('G', 1, 2)]  # 变化率列及其在 cols 中的前后下标
    docs = {}
    for s in args.source:
        label, path = s.split('=', 1)
        docs[label] = json.load(open(path, encoding='utf-8'))

    if args.template:
        wb = openpyxl.load_workbook(args.template)
        ws = wb['Sheet1']
        for r, aseg, blabel, typ, cs in ROWSHEET:
            amounts = {c: wan(amount(docs[lab], typ, cs)) for lab, c in zip(docs, cols)}
            for col, w in amounts.items():
                cell = ws[f'{col}{r}']
                cell.value = '-' if (w is None or w == 0) else w
                if w and abs(w) >= 1000:
                    cell.number_format = '#,##0.00'
            for col, i0, i1 in growth:
                v = ratio(amounts[cols[i1]], amounts[cols[i0]])
                cell = ws[f'{col}{r}']
                cell.value = '-' if v is None else v
                if v is not None:
                    cell.number_format = '0.00%'
        # 近一期表头(若模板H3是日期序列号, 统一为文本标签)
        ws['H3'] = args.period_label
        ws['H3'].font = Font(name=FONT_NAME, size=10, bold=True)
        ws['H3'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        F = lambda sz, b=False: Font(name=FONT_NAME, size=sz, bold=b)
        CA = Alignment(horizontal='center', vertical='center', wrap_text=True)
        CJ = Alignment(horizontal='justify', vertical='center', wrap_text=True)
        for r, h in ROWHEIGHTS.items():
            ws.row_dimensions[r].height = h
        for col, w in COLWIDTHS.items():
            ws.column_dimensions[col].width = w
        ws['A1'] = '近三年及近期财务数据简表'
        ws['A1'].font = F(14, True)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A2'] = '                                                  单位：万元'
        ws['A2'].font = F(14, True)
        ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
        # 表头: 前三列用年份label, 近一期列用 --period-label
        year_lbls = list(docs.keys())
        heads = ['序号', '项\u3000\u3000目', year_lbls[0], year_lbls[1], '变化率', year_lbls[2], '变化率', args.period_label]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(3, c, h)
            cell.font = F(10, True)
            cell.alignment = CA
        for r, aseg, blabel, typ, cs in ROWSHEET:
            if aseg is not None:
                a = ws.cell(r, 1, aseg); a.font = F(10, True); a.alignment = CJ
            b = ws.cell(r, 2, blabel); b.font = F(10); b.alignment = CJ
            amounts = {c: wan(amount(docs[lab], typ, cs)) for lab, c in zip(docs, cols)}
            for col, w in amounts.items():
                c = ws[f'{col}{r}']
                c.value = '-' if (w is None or w == 0) else w
                if w is not None and w != 0:
                    c.number_format = '#,##0.00' if abs(w) >= 1000 else 'General'
                c.font = F(10)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for col, i0, i1 in growth:
                v = ratio(amounts[cols[i1]], amounts[cols[i0]])
                c = ws[f'{col}{r}']
                c.value = '-' if v is None else v
                if v is not None:
                    c.number_format = '0.00%'
                c.font = F(10)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    wb.save(args.out)
    print('saved', args.out)


if __name__ == '__main__':
    main()
