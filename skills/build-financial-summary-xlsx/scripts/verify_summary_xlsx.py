#!/usr/bin/env python3
"""独立复核 build_summary_xlsx.py 产物: 逐格与源 JSON 重算比对(306 格 0 差异为目标)。
用法: python verify_summary_xlsx.py --xlsx 本部财务数据简表.xlsx --json audit2023.json audit2024.json audit2025.json audit2606.json
说明: 年份顺序=--json 传入顺序; 金额列 C/D/F/H, 变化率 E/G; '-'==0/缺失。
"""
import argparse, json, re, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
from row_defs import ROWS, LONGINV, norm_key, val, collect

_ROMAN = re.compile(r'^[一二三四五六七八九十]+、')


def lookup(doc, typ, cs):
    d = {'bs': doc['bs']['current'], 'is': doc['is']['current'], 'cf': doc['cf']['current']}[typ]
    if cs is None:
        return collect(d, *LONGINV)
    return val(d, *cs)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--json', nargs='+', required=True)
    args = ap.parse_args()
    docs = [json.load(open(p, encoding='utf-8')) for p in args.json]
    ws = openpyxl.load_workbook(args.xlsx)['Sheet1']
    cols = ['C', 'D', 'F', 'H']
    fails = 0
    for i, (label, typ, cs) in enumerate(ROWS):
        r = 4 + i
        exp = {}
        for k, col in enumerate(cols):
            v = lookup(docs[k], typ, cs)
            exp[col] = None if v is None else round(v / 1e4, 2)
        for col, y in zip(cols, range(4)):
            got = ws[f'{col}{r}'].value
            want = '-' if exp[col] in (None, 0) else exp[col]
            if got != want:
                print(f'MISMATCH r{r} {label} {col}: xlsx={got} expect={want}')
                fails += 1
        for col, i0, i1 in (('E', 0, 1), ('G', 1, 2)):
            a, b = exp[cols[i0]], exp[cols[i1]]
            if a in (None, 0) and b in (None, 0):
                want = '-'
            elif a in (None, 0):
                want = 1.0 if (b or 0) > 0 else (-1.0 if (b or 0) < 0 else 0.0)
            else:
                bv = b or 0.0
                want = round((bv - a) / abs(a), 4)
            got = ws[f'{col}{r}'].value
            if got != want:
                print(f'MISMATCH r{r} {label} {col} growth: xlsx={got} expect={want}')
                fails += 1
    print('复核差异:', fails)


if __name__ == '__main__':
    main()
