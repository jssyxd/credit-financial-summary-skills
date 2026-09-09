#!/usr/bin/env python3
"""财务分析 docx 生成模板(仿宋, A4) —— 供 Agent 按 SKILL.md 撰写正文后套版。
用法: 1) 修改下方 SECTIONS 的标题与正文(把真实分析文字填入);  2) python analysis_docx.py --out 财务分析.docx
排版: 标题黑体16居中 / 一级黑体14 / 正文仿宋12首行缩进2字符1.3倍行距 / Table Grid 表格。
"""
import argparse
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ---------- 在此填入正文(标题, 段落列表或 ('table', [[...]]) 标记) ----------
SECTIONS = [
    ('一、总体概况', [
        '（在此填写：主体定位(控股平台/购销结算中枢/经营实体)、本部 vs 合并口径的分工与资产结构特征、资产规模趋势与资本实力概况。）',
    ]),
    ('二、资产负债重点科目分析', [
        '（一）货币资金与交易性金融资产：{金额序列}；变化原因(资金投向/偿债/回款)；对冲事实(经营现金流、无受限、授信与集团统筹)；结论话术(尚可/稳健)。',
        '（二）应收账款与应收票据：{金额序列}；±30%科目重点展开(业务结构变化/结算前移/账龄)；票据清零与结算方式。',
        '（三）预付账款 /（四）其他应收款(与长投联动的债转股治理) /（五）存货 /（六）长期股权投资(核心资产占比) '
        '/（七）固定资产与在建工程 /（八）负债结构与短期借款 /（九）应付账款与预收款项 /（十）其他应付款。',
    ]),
    ('三、应收应付及存货专项分析', [
        ('table', [['科目', '2023末', '2024末', '2025末', '近一期末', '较上年', '占总资产'],
                   ['应收账款', '', '', '', '', '', ''],
                   ['预付账款', '', '', '', '', '', ''],
                   ['其他应收款', '', '', '', '', '', ''],
                   ['应付账款', '', '', '', '', '', ''],
                   ['其他应付款', '', '', '', '', '', ''],
                   ['存货', '', '', '', '', '', '']]),
        '总体判断段(良性结构 + 限定词)。',
    ]),
    ('四、损益及盈利情况分析', [
        '收入/成本/毛利率趋势；费用(财务费用随借款压降)；其他损益科目；净利润与“盈余留存子公司、近年未分红(未分配利润变动≈净利润)”的解释；正面收尾。',
    ]),
    ('五、现金流情况分析', [
        '经营(净流量序列 + 现金回收率 + 近一期表现)/投资(股权出资节奏)/筹资(主动偿债净流出) 三段。',
    ]),
    ('六、综合评价', [
        '杠杆低、结构简单、经营现金流覆盖付息债务、应收小账龄短、订单预收饱满；本部盈利体量偏小系阶段性安排；授信支持结论。',
    ]),
]


def set_font(run, name='仿宋', size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True)
    ap.add_argument('--title', default='财务分析')
    ap.add_argument('--subtitle', default='（数据区间：{区间}；除特别注明外金额单位为万元；数据来源：{来源}）')
    args = ap.parse_args()
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.6)
        sec.top_margin = sec.bottom_margin = Cm(2.8)

    def para(text, size=12, bold=False, indent=True, align=None, name='仿宋', after=6):
        p = doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Pt(size * 2)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.3
        if align:
            p.alignment = align
        r = p.add_run(text)
        set_font(r, name=name, size=size, bold=bold)
        return p

    para(args.title, 16, True, False, WD_ALIGN_PARAGRAPH.CENTER, '黑体', 10)
    para(args.subtitle, 9, False, False, WD_ALIGN_PARAGRAPH.CENTER, '楷体', 10)
    for head, body in SECTIONS:
        para(head, 14, True, False, name='黑体', after=8)
        for item in body:
            if isinstance(item, tuple) and item[0] == 'table':
                rows = item[1]
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = 'Table Grid'
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, row in enumerate(rows):
                    for j, v in enumerate(row):
                        c = t.rows[i].cells[j]
                        c.text = ''
                        r = c.paragraphs[0].add_run(str(v))
                        set_font(r, '仿宋', 9, bold=(i == 0))
                        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                para(item)
    doc.save(args.out)
    print('saved', args.out)


if __name__ == '__main__':
    main()
